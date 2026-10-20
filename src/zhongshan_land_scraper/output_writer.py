"""Output generation: JSON, CSV, Excel."""

import csv
import json
from pathlib import Path
from typing import Any

import structlog
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from zhongshan_land_scraper.models import OutputRecord

log = structlog.get_logger()

# 输出列顺序
OUTPUT_COLUMNS = [
    "成交时间",
    "竞得人",
    "用地面积",
    "成交总价",
    "宗地位置",
    "土地用途",
    "项目名称",
    "项目编号",
    "公告标题",
    "公告发布时间",
    "公告详情链接",
    "采集时间",
    "核验状态",
    "备注",
]


def _records_to_rows(records: list[OutputRecord]) -> list[list[str]]:
    """将 OutputRecord 列表转换为二维字符串列表.

    Args:
        records: 输出记录列表

    Returns:
        二维列表（含表头）
    """
    rows: list[list[str]] = [OUTPUT_COLUMNS]
    for r in records:
        row = [getattr(r, col, "") for col in OUTPUT_COLUMNS]
        rows.append(row)
    return rows


def write_json(records: list[OutputRecord], output_path: Path) -> None:
    """写入 JSON 文件.

    Args:
        records: 输出记录列表
        output_path: 输出文件路径
    """
    data: list[dict[str, Any]] = []
    for r in records:
        data.append({col: getattr(r, col, "") for col in OUTPUT_COLUMNS})

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    log.info("json_written", path=str(output_path), count=len(records))


def write_csv(records: list[OutputRecord], output_path: Path) -> None:
    """写入 CSV 文件.

    Args:
        records: 输出记录列表
        output_path: 输出文件路径
    """
    rows = _records_to_rows(records)
    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)
    log.info("csv_written", path=str(output_path), count=len(records))


def write_excel(records: list[OutputRecord], output_path: Path) -> None:
    """写入 Excel 文件.

    特性：
    - 冻结首行
    - 自动筛选
    - 合理列宽
    - 可点击详情链接
    - 保留原始值

    Args:
        records: 输出记录列表
        output_path: 输出文件路径
    """
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("Failed to create worksheet")
    ws.title = "中山市土地交易结果"

    # 写入表头
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 写入数据
    link_col_idx = OUTPUT_COLUMNS.index("公告详情链接") + 1
    for row_idx, r in enumerate(records, 2):
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, 1):
            value = getattr(r, col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            # 详情链接列设置为超链接
            if col_idx == link_col_idx and value:
                cell.hyperlink = value
                cell.value = value
                cell.style = "Hyperlink"

    # 冻结首行
    ws.freeze_panes = "A2"

    # 自动筛选
    max_row = len(records) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{max_row}"

    # 设置列宽（根据内容自适应，最小12，最大80）
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, 1):
        max_len = len(col_name)
        for row_idx in range(2, max_row + 1):
            cell_value = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, len(cell_value))
        # 限制最大宽度
        adjusted_width = min(max_len + 2, 80)
        adjusted_width = max(adjusted_width, 12)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    # 设置行高
    ws.row_dimensions[1].height = 22

    wb.save(output_path)
    log.info("excel_written", path=str(output_path), count=len(records))


def write_all_outputs(records: list[OutputRecord], output_dir: Path) -> None:
    """生成所有输出文件.

    Args:
        records: 输出记录列表
        output_dir: 输出目录
    """
    if not records:
        log.warning("no_records_to_output")
        return

    output_dir.mkdir(parents=True, exist_ok=True)

    write_json(records, output_dir / "zhongshan_land_transactions.json")
    write_csv(records, output_dir / "zhongshan_land_transactions.csv")
    write_excel(records, output_dir / "zhongshan_land_transactions.xlsx")

    log.info("all_outputs_generated", dir=str(output_dir), count=len(records))
